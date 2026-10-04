"""Reading job records out of an arbitrary Excel workbook.

The workbook is inspected at runtime. Nothing about its shape is assumed: not
the sheet name, not the header row position, not the column names. Two
different job boards export wildly different headers for the same field, and
hard-coding one export's names would break silently on the next.

Detection is confidence-scored rather than best-guess. When two columns are
equally plausible for a required field the module refuses to choose and hands
the decision back to the caller, because a silent wrong guess produces a
hundred confidently wrong emails.
"""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any, Iterable, Optional

from openpyxl import load_workbook

# How many rows to look through for the header. Exports often carry a title
# row, a blank row, or a filter banner before the real headers.
_HEADER_SEARCH_ROWS = 12

# A header row is only accepted if it has at least this many non-empty cells.
_MIN_HEADER_CELLS = 2

# Detection vocabulary. Each field lists phrases that would appear in a header
# for it, strongest first. These are *hints for scoring*, not a whitelist -
# an unrecognised header is still offered to the user for manual mapping.
_FIELD_HINTS: dict[str, dict[str, Any]] = {
    "role": {
        "label": "Job Role / Title",
        "required": True,
        "exact": ("job title", "job role", "role title", "position title",
                  "role", "title", "position", "job", "designation", "vacancy"),
        "contains": ("job title", "job role", "position", "designation", "role", "title"),
        "avoid": ("company", "description", "url", "link", "id", "type", "level"),
    },
    "description": {
        "label": "Job Description",
        "required": True,
        "exact": ("job description", "description", "jd", "job details",
                  "details", "requirements", "job summary", "summary",
                  "job post", "posting", "job text", "full description"),
        "contains": ("description", "requirement", "responsib", "details",
                     "summary", "job post", "jd"),
        "avoid": ("company", "title", "url", "link", "short"),
    },
    "company": {
        "label": "Company",
        "required": False,
        "exact": ("company", "company name", "employer", "organisation",
                  "organization", "client", "firm", "account"),
        "contains": ("company", "employer", "organis", "organiz", "client"),
        "avoid": ("description", "size", "url", "link", "id"),
    },
    "recipient_email": {
        "label": "Recipient Email",
        "required": False,
        "exact": ("email", "contact email", "recruiter email", "hr email",
                  "hiring manager email", "email address", "contact",
                  "recipient", "recipient email"),
        "contains": ("email", "e-mail", "mail"),
        "avoid": ("company", "description", "url", "domain"),
    },
    "location": {
        "label": "Location",
        "required": False,
        "exact": ("location", "job location", "city", "region", "country",
                  "based in", "office location", "work location"),
        "contains": ("location", "city", "country", "region"),
        "avoid": ("description", "url"),
    },
    "remote": {
        "label": "Remote / Work Mode",
        "required": False,
        "exact": ("remote", "work mode", "workplace", "work type",
                  "remote status", "onsite", "hybrid", "arrangement"),
        "contains": ("remote", "work mode", "workplace", "hybrid", "onsite"),
        "avoid": ("description",),
    },
    "experience": {
        "label": "Experience",
        "required": False,
        "exact": ("experience", "years of experience", "experience required",
                  "min experience", "seniority", "experience level"),
        "contains": ("experience", "seniority"),
        "avoid": ("description",),
    },
    "skills": {
        "label": "Skills",
        "required": False,
        "exact": ("skills", "required skills", "key skills", "tech stack",
                  "technologies", "keywords", "tags"),
        "contains": ("skill", "technolog", "tech stack", "keyword", "tag"),
        "avoid": ("description",),
    },
    "hiring_manager": {
        "label": "Hiring Manager",
        "required": False,
        "exact": ("hiring manager", "recruiter", "contact name", "contact person",
                  "hr", "poc", "point of contact"),
        "contains": ("hiring manager", "recruiter", "contact name", "contact person"),
        "avoid": ("email", "description"),
    },
    "url": {
        "label": "Job URL",
        "required": False,
        "exact": ("url", "link", "job url", "job link", "apply link",
                  "posting url", "source"),
        "contains": ("url", "link"),
        "avoid": ("description",),
    },
}

REQUIRED_FIELDS = tuple(name for name, spec in _FIELD_HINTS.items() if spec["required"])

# Below this, a match is not trusted at all. Within _AMBIGUITY_MARGIN of the
# runner-up, the field is reported ambiguous rather than picked.
_MIN_CONFIDENCE = 40
_AMBIGUITY_MARGIN = 15

# Added when a column's header and its contents both point at the same field.
# Agreement is much stronger evidence than either signal alone.
_AGREEMENT_BONUS = 25

_EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[A-Za-z]{2,}$")
_URL_RE = re.compile(r"^https?://", re.I)
_ISO_DATE_RE = re.compile(r"^\d{4}-\d{2}-\d{2}[T ]")
_NUMERIC_RE = re.compile(r"^-?\d+(\.\d+)?$")
_BOOLEAN_VALUES = {"true", "false", "yes", "no", "none", "only", "y", "n", "0", "1"}

# Words that make a short text value look like a job title.
_ROLE_WORDS = (
    "engineer", "developer", "manager", "analyst", "scientist", "designer",
    "architect", "consultant", "specialist", "lead", "director", "intern",
    "administrator", "technician", "researcher", "officer", "associate",
    "coordinator", "executive", "head of", "founding", "staff", "senior",
    "junior", "principal", "full-stack", "fullstack", "frontend", "backend",
)


def _looks_like_header_row(values: Iterable[Any]) -> bool:
    """Whether a row is headers rather than the first record.

    Scrapes and API exports frequently ship with no header row at all. Reading
    the first record as headers silently loses a job AND destroys every
    mapping, so the row has to earn its status: header cells are short, mostly
    unique, and none of them look like a description, a URL, a timestamp or a
    flag.
    """
    cells = [str(value).strip() for value in values if str(value or "").strip()]
    if len(cells) < _MIN_HEADER_CELLS:
        return False

    for cell in cells:
        lowered = cell.lower()
        if len(cell) > 60:
            return False                      # a description, not a header
        if _URL_RE.match(cell) or _ISO_DATE_RE.match(cell):
            return False
        if lowered in _BOOLEAN_VALUES:
            return False                      # a flag value
        if _EMAIL_RE.match(cell):
            return False

    numeric = sum(1 for cell in cells if _NUMERIC_RE.match(cell))
    if numeric > len(cells) * 0.25:
        return False                          # ids and counts, not names

    if len(set(cell.lower() for cell in cells)) < len(cells) * 0.9:
        return False                          # headers are unique

    # Finally: a genuine header row names at least one field we recognise.
    return any(
        _score_header(_normalize_header(cell), spec) >= _MIN_CONFIDENCE
        for cell in cells
        for spec in _FIELD_HINTS.values()
    )


def _column_letter(index: int) -> str:
    """A, B, ... Z, AA - for naming columns in a headerless workbook."""
    letters = ""
    index += 1
    while index:
        index, remainder = divmod(index - 1, 26)
        letters = chr(65 + remainder) + letters
    return letters


def _normalize_header(value: Any) -> str:
    """Lower-cased header with punctuation and separators flattened."""
    text = str(value or "").strip().lower()
    text = re.sub(r"[_\-./\\]+", " ", text)
    text = re.sub(r"[^a-z0-9\s&+#]", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def _score_header(header: str, spec: dict[str, Any]) -> int:
    """How strongly a header matches a field, 0-100."""
    if not header:
        return 0

    for avoid in spec["avoid"]:
        if avoid in header and not any(header == exact for exact in spec["exact"]):
            return 0

    if header in spec["exact"]:
        # Earlier entries are stronger: "job title" beats a bare "title".
        position = spec["exact"].index(header)
        return max(70, 100 - position * 3)

    for phrase in spec["contains"]:
        if phrase in header:
            # A header that is mostly the phrase is a better match than one
            # that merely contains it inside a longer name.
            coverage = len(phrase) / max(len(header), 1)
            return int(40 + 40 * min(coverage, 1.0))

    return 0


def inspect_workbook(path: str | Path, sheet_name: Optional[str] = None) -> dict[str, Any]:
    """Read a workbook's structure without committing to any interpretation.

    Returns the sheet names, the header row that was found, the columns, a few
    sample values per column, and the row count. The caller decides what to do
    with it.
    """
    workbook = load_workbook(filename=str(path), read_only=True, data_only=True)
    try:
        sheet_names = list(workbook.sheetnames)
        if not sheet_names:
            raise ValueError("The workbook contains no worksheets.")

        active_name = sheet_name if sheet_name in sheet_names else sheet_names[0]
        sheet = workbook[active_name]

        rows = []
        for index, row in enumerate(sheet.iter_rows(values_only=True), start=1):
            rows.append(row)
            if index >= _HEADER_SEARCH_ROWS:
                break

        header_row_index, headers = _find_header_row(rows)
        has_headers = bool(headers)

        if not has_headers:
            # Every row is data. Columns are named by position and identified
            # from their contents instead.
            width = sheet.max_column or (max((len(row) for row in rows), default=0))
            if not width:
                raise ValueError(f"Sheet '{active_name}' appears to be empty.")
            headers = [f"Column {_column_letter(position)}" for position in range(width)]

        samples, stats = _collect_samples(sheet, header_row_index, len(headers))
        data_rows = max(0, (sheet.max_row or 0) - header_row_index)

        columns = [
            {
                "index": position,
                "name": str(name),
                "normalized": "" if not has_headers else _normalize_header(name),
                "samples": samples.get(position, []),
                "stats": stats.get(position, {"avg_length": 0, "max_length": 0, "values": []}),
            }
            for position, name in enumerate(headers)
            # A headerless workbook keeps empty columns so positions stay
            # aligned; a headed one drops unnamed columns.
            if not has_headers or str(name or "").strip()
        ]

        return {
            "path": str(path),
            "sheet_names": sheet_names,
            "sheet": active_name,
            "header_row": header_row_index,
            "has_headers": has_headers,
            "columns": columns,
            "row_count": data_rows,
        }
    finally:
        workbook.close()


def _find_header_row(rows: list[tuple]) -> tuple[int, list[Any]]:
    """The first row that looks like headers rather than a title or a blank.

    Exports frequently open with a report title, an export timestamp or a
    blank row, so row 1 cannot be assumed to hold the headers.
    """
    for index, row in enumerate(rows, start=1):
        if not row:
            continue
        if _looks_like_header_row(row):
            return index, list(row)

    # No header row anywhere. Signalled with 0 so the caller knows every row
    # is data and columns must be identified from their values.
    return 0, []


def _collect_samples(
    sheet, header_row_index: int, width: int, limit: int = 3
) -> tuple[dict[int, list[str]], dict[int, dict[str, Any]]]:
    """Display samples, plus statistics measured on the untruncated values.

    The samples are clipped so the mapping UI stays readable, but scoring must
    see the real lengths - a 5000-character job description clipped to 120
    looks like a job title.
    """
    samples: dict[int, list[str]] = {}
    stats: dict[int, dict[str, Any]] = {}
    seen_rows = 0

    for index, row in enumerate(sheet.iter_rows(values_only=True), start=1):
        if index <= header_row_index:
            continue
        for position in range(min(width, len(row))):
            value = str(row[position] or "").strip()
            if not value:
                continue
            if len(samples.setdefault(position, [])) < limit:
                samples[position].append(value[:120])
            entry = stats.setdefault(position, {"lengths": [], "values": []})
            entry["lengths"].append(len(value))
            if len(entry["values"]) < 12:
                entry["values"].append(value[:200])
        seen_rows += 1
        if seen_rows >= 25:
            break

    for entry in stats.values():
        lengths = entry.pop("lengths")
        entry["avg_length"] = sum(lengths) / len(lengths) if lengths else 0
        entry["max_length"] = max(lengths) if lengths else 0

    return samples, stats


def _looks_like_person(value: str) -> bool:
    """Two to four capitalised words, no digits or punctuation - a name."""
    parts = value.split()
    return (
        1 < len(parts) <= 4
        and all(part[:1].isupper() for part in parts)
        and not any(char.isdigit() for char in value)
        and not any(char in value for char in ";,/@")
    )


def _score_values(field: str, column: dict[str, Any]) -> int:
    """Identify a column from what it contains, not what it is called.

    This is what makes a headerless export usable: a column of 900-character
    paragraphs is the job description whatever it is named, and a column of
    "Senior Backend Engineer" strings is the role.
    """
    stats = column.get("stats") or {}
    samples = [value for value in (stats.get("values") or column.get("samples") or []) if value.strip()]
    if not samples:
        return 0

    average = stats.get("avg_length") or (sum(len(v) for v in samples) / len(samples))

    def share(predicate) -> float:
        return sum(1 for value in samples if predicate(value)) / len(samples)

    if field == "description":
        # Long prose, and not a URL.
        if average >= 200 and share(lambda v: _URL_RE.match(v)) < 0.5:
            return 90
        if average >= 120:
            return 55
        return 0

    if field == "role":
        if average > 80 or average < 4:
            return 0
        if share(lambda v: any(word in v.lower() for word in _ROLE_WORDS)) >= 0.5:
            return 88
        return 0

    if field == "recipient_email":
        return 85 if share(lambda v: _EMAIL_RE.match(v.strip())) >= 0.6 else 0

    if field == "url":
        return 80 if share(lambda v: bool(_URL_RE.match(v))) >= 0.6 else 0

    if field == "location":
        # "New York, NY, US" or "San Francisco, CA, US"
        if average > 90:
            return 0
        if share(lambda v: "," in v and not _URL_RE.match(v)) >= 0.6 and average < 60:
            return 60
        return 0

    if field == "skills":
        # "Node.js; Python; React; TypeScript"
        if share(lambda v: v.count(";") >= 2 or v.count(",") >= 3) >= 0.5 and 15 < average < 200:
            return 70
        return 0

    if field == "remote":
        vocabulary = {"remote", "onsite", "on-site", "hybrid", "only", "no", "yes", "none"}
        return 65 if share(lambda v: v.strip().lower() in vocabulary) >= 0.8 else 0

    if field == "company":
        if not (2 < average < 45):
            return 0
        if share(lambda v: bool(_URL_RE.match(v)) or bool(_NUMERIC_RE.match(v))) > 0:
            return 0
        if share(lambda v: v.strip().lower() in _BOOLEAN_VALUES) > 0:
            return 0
        # A location list ("New York, NY, US; New York") is not a company.
        if share(lambda v: ";" in v or v.count(",") >= 2) >= 0.5:
            return 0
        # Nor is a person's name - that is the hiring manager column.
        if share(_looks_like_person) >= 0.6:
            return 0
        if share(lambda v: v[:1].isupper()) >= 0.8:
            return 70
        return 0

    if field == "hiring_manager":
        return 75 if share(_looks_like_person) >= 0.6 else 0

    return 0


def detect_columns(columns: list[dict[str, Any]]) -> dict[str, Any]:
    """Map fields to columns, refusing to guess when the choice is unclear.

    Returns the confident mappings, the fields that were ambiguous (with their
    candidates), and the required fields that could not be found at all.
    """
    mapping: dict[str, int] = {}
    confidence: dict[str, int] = {}
    ambiguous: dict[str, list[dict[str, Any]]] = {}
    missing: list[str] = []
    taken: set[int] = set()

    for field, spec in _FIELD_HINTS.items():
        scored = []
        for column in columns:
            # Header and values are scored independently, then combined. A
            # column that satisfies BOTH outranks one that satisfies only
            # the header - which is what separates a column headed "role"
            # holding category codes ("eng", "ops") from a column headed
            # "title" holding actual job titles. The header alone would pick
            # the wrong one.
            header_score = _score_header(column["normalized"], spec)
            value_score = _score_values(field, column)
            score = max(header_score, value_score)
            if header_score > 0 and value_score > 0:
                # Deliberately uncapped: capping at 100 collapsed the gap
                # between a column agreeing on both signals and one matching
                # only its header, leaving them falsely ambiguous.
                score += _AGREEMENT_BONUS
            if score > 0 and column["index"] not in taken:
                scored.append((score, column))

        scored.sort(key=lambda item: (-item[0], item[1]["index"]))

        if not scored or scored[0][0] < _MIN_CONFIDENCE:
            if spec["required"]:
                missing.append(field)
            continue

        best_score, best_column = scored[0]
        runner_up = scored[1][0] if len(scored) > 1 else 0

        if best_score - runner_up < _AMBIGUITY_MARGIN and runner_up >= _MIN_CONFIDENCE:
            # Two equally plausible columns. Refuse rather than coin-flip.
            ambiguous[field] = [
                {"index": column["index"], "name": column["name"], "score": score}
                for score, column in scored[:4]
            ]
            if spec["required"]:
                missing.append(field)
            continue

        mapping[field] = best_column["index"]
        confidence[field] = min(100, best_score)
        taken.add(best_column["index"])

    return {
        "mapping": mapping,
        "confidence": confidence,
        "ambiguous": ambiguous,
        "missing_required": [field for field in missing if _FIELD_HINTS[field]["required"]],
        "labels": {field: spec["label"] for field, spec in _FIELD_HINTS.items()},
        "required_fields": list(REQUIRED_FIELDS),
    }


def _email_column_bonus(column: dict[str, Any]) -> int:
    """A column whose values are email addresses, whatever its header says."""
    samples = column.get("samples") or []
    if not samples:
        return 0
    hits = sum(1 for value in samples if _EMAIL_RE.match(value.strip()))
    return 85 if hits and hits >= len(samples) * 0.6 else 0


def read_jobs(
    path: str | Path,
    mapping: dict[str, int],
    *,
    sheet: Optional[str] = None,
    header_row: int = 1,
    limit: Optional[int] = None,
) -> list[dict[str, Any]]:
    """Read job rows using an agreed mapping.

    A row missing a role or a description is returned with a `skip_reason`
    rather than dropped, so the batch report can account for every row.
    """
    workbook = load_workbook(filename=str(path), read_only=True, data_only=True)
    try:
        worksheet = workbook[sheet] if sheet in workbook.sheetnames else workbook[workbook.sheetnames[0]]
        jobs: list[dict[str, Any]] = []

        for index, row in enumerate(worksheet.iter_rows(values_only=True), start=1):
            if index <= header_row:
                continue
            if limit is not None and len(jobs) >= limit:
                break

            record: dict[str, Any] = {"row": index}
            for field, position in mapping.items():
                value = row[position] if position < len(row) else None
                record[field] = str(value).strip() if value is not None else ""

            if not any(str(value).strip() for key, value in record.items() if key != "row"):
                continue  # entirely blank row

            role = record.get("role", "")
            description = record.get("description", "")
            if not role.strip():
                record["skip_reason"] = "No job role in this row."
            elif len(description.strip()) < 40:
                record["skip_reason"] = "Job description missing or too short to analyse."

            jobs.append(record)

        return jobs
    finally:
        workbook.close()


def is_valid_email(value: str) -> bool:
    return bool(_EMAIL_RE.match((value or "").strip()))
